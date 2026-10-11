import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from apps.curriculum.models import Student, StudentClass, ClassEnrollment
from apps.journals.models import Journal, Lesson, Grade, SummaryGrade, Attendance, SummaryGradeType
from apps.journals.services.scheduling import SUMMARY_TYPES
from apps.journals.services.assessment import legacy_summary_groups


class ExportService:
    @staticmethod
    def export_students(students, student_class=None, file_format='xlsx'):
        """Експортує список учнів у форматі XLSX або CSV (сумісному для повторного імпорту)."""
        headers = ["Код/ID учня", "Прізвище", "Ім'я", "По батькові", "Клас", "Статус (ACTIVE/TRANSFERRED/EXPELLED)"]
        rows = []
        for s in students:
            # Визначаємо клас
            c_name = student_class.name if student_class else ""
            if not c_name:
                active_enr = s.enrollments.filter(is_active=True).first()
                if active_enr:
                    c_name = active_enr.student_class.name
            rows.append([
                s.internal_id or "",
                s.last_name,
                s.first_name,
                s.middle_name or "",
                c_name,
                s.status
            ])

        if file_format == 'csv':
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(headers)
            for r in rows:
                writer.writerow(r)
            return output.getvalue().encode('utf-8-sig'), 'text/csv'
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Список учнів"

            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            ws.append(headers)
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            for r in rows:
                ws.append(r)
                row_idx = ws.max_row
                for col_idx in range(1, len(headers) + 1):
                    ws.cell(row=row_idx, column=col_idx).border = thin_border

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    @staticmethod
    def export_topics(journal, semester=None, file_format='xlsx'):
        """Export every lesson and unlinked plan topic with stable journal-scoped IDs."""
        from .topic_service import TopicImportService

        plans = list(journal.plan_topics.order_by('topic_number', 'pk'))
        reserved = {p.topic_number for p in plans}
        lessons = journal.lessons.select_related('linked_plan_topic').order_by('date', 'lesson_number', 'pk')
        linked_ids = set()
        rows = []
        next_number = 1
        for lesson in lessons:
            if lesson.linked_plan_topic_id:
                number = lesson.linked_plan_topic.topic_number
                linked_ids.add(lesson.linked_plan_topic_id)
            else:
                while next_number in reserved:
                    next_number += 1
                number = next_number
                reserved.add(number)
                next_number += 1
            if semester and lesson.semester_id != semester.pk:
                continue
            rows.append([number, lesson.topic, lesson.date.isoformat(), lesson.homework,
                         lesson.lesson_type, lesson.notes, lesson.pk, lesson.lesson_number,
                         lesson.linked_plan_topic_id or '', 'LESSON', journal.pk])
        for plan in plans:
            if plan.pk in linked_ids:
                continue
            if semester and plan.planned_date and not semester.start_date <= plan.planned_date <= semester.end_date:
                continue
            rows.append([plan.topic_number, plan.title, plan.planned_date.isoformat() if plan.planned_date else '',
                         plan.homework, plan.lesson_type, plan.notes, '', '', plan.pk, 'PLAN', journal.pk])
        return TopicImportService.render_file(rows, file_format)

    @staticmethod
    def export_journal_xlsx(journal, semester=None, start_date=None, end_date=None):
        """
        Генерує повноцінний файл журналу в Excel:
        - Шапка: заклад, клас, предмет, вчитель, навчальний рік
        - Сітка: учні у рядках, уроки у стовпцях, оцінки/відвідування у клітинках
        - Підсумкові оцінки (Тематичні, Семестрові, Річна)
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Журнал {journal.student_class.name}"

        # 1. Верхній блок метаданих
        title_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        meta_label_font = Font(name="Arial", size=10, bold=True, color="475569")
        meta_value_font = Font(name="Arial", size=10, color="0F172A")

        ws.merge_cells('A1:K1')
        ws['A1'] = journal.school.name
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal="center")

        ws.merge_cells('A2:K2')
        ws['A2'] = f"ЕЛЕКТРОННИЙ ЖУРНАЛ: {journal.student_class.name} КЛАС — ПРЕДМЕТ «{journal.subject.name.upper()}»"
        ws['A2'].font = Font(name="Arial", size=12, bold=True)
        ws['A2'].alignment = Alignment(horizontal="center")

        ws['A3'] = "Вчитель:"
        ws['A3'].font = meta_label_font
        ws['B3'] = journal.teacher.get_full_name_ukr()
        ws['B3'].font = meta_value_font

        ws['D3'] = "Навчальний рік:"
        ws['D3'].font = meta_label_font
        ws['E3'] = journal.academic_year.name
        ws['E3'].font = meta_value_font

        period_label = semester.name if semester else "Весь навчальний рік"
        ws['G3'] = "Період:"
        ws['G3'].font = meta_label_font
        ws['H3'] = period_label
        ws['H3'].font = meta_value_font

        # 2. Отримуємо уроки
        lessons_qs = Lesson.objects.filter(journal=journal).order_by('date', 'lesson_number')
        if semester:
            lessons_qs = lessons_qs.filter(semester=semester)
        if start_date:
            lessons_qs = lessons_qs.filter(date__gte=start_date)
        if end_date:
            lessons_qs = lessons_qs.filter(date__lte=end_date)
        lessons = list(lessons_qs)

        # 3. Отримуємо учнів класу (включаючи історичних, що мали оцінки у цьому журналі)
        enrolled_student_ids = ClassEnrollment.objects.filter(
            student_class=journal.student_class
        ).values_list('student_id', flat=True)

        graded_student_ids = Grade.objects.filter(
            lesson__journal=journal
        ).values_list('student_id', flat=True)

        all_student_ids = set(enrolled_student_ids).union(set(graded_student_ids))
        students = list(Student.objects.filter(id__in=all_student_ids).order_by('last_name', 'first_name'))

        # Підготовка матриць оцінок та відвідування
        grades_map = {} # (lesson_id, student_id) -> value
        for g in Grade.objects.filter(lesson__in=lessons, is_approved=True):
            grades_map[(g.lesson_id, g.student_id)] = g.value

        attendance_map = {} # (lesson_id, student_id) -> symbol
        for a in Attendance.objects.filter(lesson__in=lessons):
            symbol = a.get_symbol_ukr()
            if symbol:
                attendance_map[(a.lesson_id, a.student_id)] = symbol

        summary_grades_map = {}
        legacy_columns = []
        for key, group in legacy_summary_groups(journal).items():
            if any(sg.value for sg in group['grades'].values()):
                legacy_columns.append((group['label'], key))
                summary_grades_map.update({(sid, key): sg.value for sid, sg in group['grades'].items()})

        # 4. Формуємо шапку таблиці (починаючи з рядка 5)
        # Рядок 5: Номери уроків або типи
        # Рядок 6: Дати уроків
        header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        th_font = Font(name="Arial", size=9, bold=True, color="1E293B")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        double_bottom = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='medium', color='1E3A8A')
        )

        ws['A5'] = "№"
        ws['A6'] = "з/п"
        ws['B5'] = "Прізвище, ім'я учня"
        ws['B6'] = "(ПІБ)"

        col_idx = 3
        for l in lessons:
            col_letter = get_column_letter(col_idx)
            ws.cell(row=5, column=col_idx, value=l.get_lesson_type_display() if l.lesson_type != 'ORDINARY' else f"Урок {l.lesson_number}")
            ws.cell(row=6, column=col_idx, value=l.date.strftime('%d.%m'))
            ws.column_dimensions[col_letter].width = 6
            col_idx += 1

        # Summary columns
        sum_cols = legacy_columns + [("Сер. бал", "AVG")]
        sum_col_indices = {}
        for label, stype in sum_cols:
            col_letter = get_column_letter(col_idx)
            ws.cell(row=5, column=col_idx, value="Підсумок")
            ws.cell(row=6, column=col_idx, value=label)
            ws.column_dimensions[col_letter].width = 8
            sum_col_indices[stype] = col_idx
            col_idx += 1

        # Apply header styling
        for r in [5, 6]:
            for c in range(1, col_idx):
                cell = ws.cell(row=r, column=c)
                cell.fill = header_fill
                cell.font = th_font
                cell.alignment = center_align
                cell.border = double_bottom if r == 6 else thin_border

        ws.column_dimensions['A'].width = 5
        ws.column_dimensions['B'].width = 28

        # 5. Заповнення рядків учнів
        current_row = 7
        for idx, student in enumerate(students, start=1):
            ws.cell(row=current_row, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=2, value=student.get_full_name_ukr()).alignment = Alignment(horizontal="left")

            numeric_grades = []
            for l_idx, lesson in enumerate(lessons, start=3):
                cell = ws.cell(row=current_row, column=l_idx)
                # Перевіряємо чи є оцінка або відвідуваність
                val = grades_map.get((lesson.id, student.id), "")
                att = attendance_map.get((lesson.id, student.id), "")

                display_val = val
                if att and not val:
                    display_val = att
                elif att and val:
                    display_val = f"{val}/{att}"

                cell.value = display_val
                cell.alignment = center_align
                cell.border = thin_border

                if val and val.isdigit() and lesson.lesson_type not in SUMMARY_TYPES:
                    numeric_grades.append(int(val))

            # Підсумкові оцінки
            for stype, col_num in sum_col_indices.items():
                cell = ws.cell(row=current_row, column=col_num)
                cell.alignment = center_align
                cell.border = thin_border
                if stype == "AVG":
                    if numeric_grades:
                        avg = round(sum(numeric_grades) / len(numeric_grades), 1)
                        cell.value = avg
                        cell.font = Font(name="Arial", size=9, italic=True, color="64748B")
                    else:
                        cell.value = "-"
                else:
                    cell.value = summary_grades_map.get((student.id, stype), "")
                    cell.font = Font(name="Arial", size=9, bold=True, color="1E3A8A")

            # Border for row headers
            ws.cell(row=current_row, column=1).border = thin_border
            ws.cell(row=current_row, column=2).border = thin_border

            current_row += 1

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
