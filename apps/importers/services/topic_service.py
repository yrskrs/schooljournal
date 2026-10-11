import csv
import datetime
import io
import re
from collections import Counter, defaultdict

from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from django.db import transaction

from apps.journals.models import Journal, LessonPlanTopic, Lesson, LessonType, AuditLog
from apps.journals.services.scheduling import SUMMARY_TYPES, next_lesson_number


class TopicImportService:
    HEADERS = [
        '№ уроку', 'Тема уроку', 'Дата уроку (РРРР-ММ-ДД або ДД.ММ.РРРР)',
        'Домашнє завдання', 'Тип уроку', 'Примітки', 'ID уроку',
        'Номер уроку на дату', 'ID теми плану', 'Вид запису (LESSON/PLAN)', 'ID журналу',
    ]
    MODES = {'add_and_update', 'add_only', 'update_only'}
    TARGETS = {'lessons_and_plan', 'plan_only'}

    @classmethod
    def generate_template(cls, file_format='xlsx'):
        """An empty template: example dates must never become real lessons."""
        return cls.render_file([], file_format)

    @classmethod
    def render_file(cls, rows, file_format='xlsx'):
        if file_format == 'csv':
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(cls.HEADERS)
            writer.writerows(rows)
            return output.getvalue().encode('utf-8-sig'), 'text/csv'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Тематичний план'
        ws.append(cls.HEADERS)
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='1E3A8A')
            cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for row in rows:
            ws.append(row)
        for col in ws.columns:
            width = max(len(str(cell.value or '')) for cell in col) + 4
            ws.column_dimensions[col[0].column_letter].width = min(max(width, 15), 65)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    @classmethod
    def _parse_html_table(cls, content):
        if not content:
            return []
        if isinstance(content, str):
            decoded = content
        else:
            for encoding in ('utf-8-sig', 'utf-8', 'cp1251', 'windows-1251'):
                try:
                    decoded = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                decoded = content.decode('utf-8', errors='ignore')
        soup = BeautifulSoup(decoded, 'html.parser')
        table = soup.find('table')
        if not table:
            return []
        raw_rows = []
        for tr in table.find_all('tr'):
            cells = tr.find_all(['th', 'td'])
            if cells:
                raw_rows.append([c.get_text(' ', strip=True) for c in cells])
        return raw_rows

    @classmethod
    def parse_file(cls, file_obj, filename):
        ext = filename.lower()
        if ext.endswith('.csv'):
            content = file_obj.read()
            for encoding in ('utf-8-sig', 'cp1251'):
                try:
                    decoded = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ValueError('Не вдалося розпізнати кодування CSV. Використовуйте UTF-8.')
            try:
                dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            raw_rows = list(csv.reader(io.StringIO(decoded), dialect))
        elif ext.endswith(('.xlsx', '.xls', '.html', '.htm')):
            raw_rows = []
            if ext.endswith('.xlsx'):
                try:
                    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
                    try:
                        raw_rows = list(wb.active.iter_rows(values_only=True))
                    finally:
                        wb.close()
                except Exception:
                    file_obj.seek(0)
                    raw_rows = cls._parse_html_table(file_obj.read())
            else:
                content = file_obj.read()
                raw_rows = cls._parse_html_table(content)
                if not raw_rows:
                    raise ValueError('Не вдалося розпізнати таблицю у файлі XLS/HTML.')
        else:
            raise ValueError('Підтримуються лише XLSX, XLS та CSV.')
        if not raw_rows:
            return []
        header_map = cls._map_headers([str(c or '').strip() for c in raw_rows[0]])
        rows = []
        for row_num, raw_row in enumerate(raw_rows[1:], start=2):
            if not any(c is not None and str(c).strip() for c in raw_row):
                continue
            item = {'_row_num': row_num}
            for key, index in header_map.items():
                value = raw_row[index] if index is not None and index < len(raw_row) else None
                if isinstance(value, (datetime.datetime, datetime.date)):
                    value = value.strftime('%Y-%m-%d')
                item[key] = str(value).strip() if value is not None else ''
            rows.append(item)
        return rows

    @staticmethod
    def _map_headers(headers):
        mapping = dict.fromkeys(('topic_number', 'title', 'date', 'homework', 'lesson_type',
                                 'notes', 'lesson_id', 'day_lesson_number', 'plan_topic_id',
                                 'record_kind', 'journal_id'))
        for index, header in enumerate(headers):
            h = header.lower().strip()
            if h in ('id уроку', 'lesson_id'):
                key = 'lesson_id'
            elif h in ('id теми плану', 'plan_topic_id'):
                key = 'plan_topic_id'
            elif h in ('id журналу', 'journal_id'):
                key = 'journal_id'
            elif h == 'day_lesson_number' or 'на дату' in h:
                key = 'day_lesson_number'
            elif h == 'record_kind' or 'вид запису' in h:
                key = 'record_kind'
            elif any(k in h for k in ('№', 'номер', 'порядок', 'number')):
                key = 'topic_number'
            elif any(k in h for k in ('тип', 'type')):
                key = 'lesson_type'
            elif any(k in h for k in ('тема', 'назва', 'topic', 'title', 'зміст', 'content')):
                key = 'title'
            elif any(k in h for k in ('дата', 'число', 'date')):
                key = 'date'
            elif any(k in h for k in ('дз', 'домашнє', 'завдання', 'homework')):
                key = 'homework'
            elif any(k in h for k in ('приміт', 'нотат', 'notes', 'заміна', 'substitut')):
                key = 'notes'
            else:
                continue
            mapping[key] = index
        if mapping['title'] is None:
            raise ValueError('У файлі немає колонки «Тема уроку» (або «Зміст»). Завантажте шаблон.')
        return mapping

    @staticmethod
    def _parse_date(value, base_year=None):
        if not value:
            return None
        value = str(value).strip()
        for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%d/%m/%Y', '%Y.%m.%d', '%d-%m-%Y'):
            try:
                return datetime.datetime.strptime(value, fmt).date()
            except ValueError:
                pass
        ukr_months = {
            'січ': 1, 'лют': 2, 'бер': 3, 'квіт': 4, 'трав': 5, 'черв': 6,
            'лип': 7, 'серп': 8, 'вер': 9, 'жовт': 10, 'лист': 11, 'груд': 12
        }
        wm = re.search(r'(\d{1,2})\s+([а-яА-Яіїєґ]+)(?:\s+(\d{4}))?', value.lower())
        if wm:
            day = int(wm.group(1))
            m_txt = wm.group(2)
            mo = next((m for pfx, m in ukr_months.items() if m_txt.startswith(pfx)), None)
            if mo:
                yr = int(wm.group(3)) if wm.group(3) else (base_year or datetime.date.today().year)
                if not wm.group(3) and base_year:
                    yr = base_year if mo >= 8 else base_year + 1
                try:
                    return datetime.date(yr, mo, day)
                except ValueError:
                    pass
        dm = re.search(r'^(\d{1,2})[./](\d{1,2})(?:[./](\d{2,4}))?$', value)
        if dm and base_year:
            day, mo = int(dm.group(1)), int(dm.group(2))
            yr = int(dm.group(3)) if dm.group(3) else (base_year if mo >= 8 else base_year + 1)
            if yr < 100:
                yr += 2000
            try:
                return datetime.date(yr, mo, day)
            except ValueError:
                pass
        try:
            return datetime.date.fromisoformat(value[:10])
        except ValueError:
            return None

    @classmethod
    def validate_and_preview(cls, rows, journal, target='lessons_and_plan'):
        if target not in cls.TARGETS:
            raise ValueError('Невідомий спосіб імпорту.')
        plans = list(journal.plan_topics.all())
        lessons = list(journal.lessons.all())
        plans_by_id = {p.pk: p for p in plans}
        lessons_by_id = {l.pk: l for l in lessons}
        plans_by_number, lessons_by_date, lessons_by_plan = defaultdict(list), defaultdict(list), defaultdict(list)
        for p in plans:
            plans_by_number[p.topic_number].append(p)
        for lesson in lessons:
            lessons_by_date[lesson.date].append(lesson)
            if lesson.linked_plan_topic_id:
                lessons_by_plan[lesson.linked_plan_topic_id].append(lesson)
        preview = []
        for index, raw in enumerate(rows, 1):
            item = {key: str(raw.get(key) or '').strip() for key in cls._map_headers(cls.HEADERS)}
            item['row_num'] = raw.get('_row_num', raw.get('row_num', index + 1))
            errors = []
            if not item['title'] or len(item['title']) > 255:
                errors.append('Тема обов’язкова та має містити до 255 символів.')
            if len(item['notes']) > 255:
                errors.append('Примітка має містити до 255 символів.')
            for key, maximum in (('topic_number', 2147483647), ('day_lesson_number', 32767),
                                 ('lesson_id', None), ('plan_topic_id', None), ('journal_id', None)):
                try:
                    value = int(item[key]) if item[key] else (index if key == 'topic_number' else None)
                    if value is not None and (value < 1 or (maximum and value > maximum)):
                        raise ValueError()
                    item[key] = value
                except ValueError:
                    item[key] = None
            if item['journal_id'] and item['journal_id'] != journal.pk:
                errors.append('Файл належить іншому журналу. Завантажте дані цього журналу.')
            base_year = journal.academic_year.start_date.year if journal.academic_year else None
            lesson_date = cls._parse_date(item['date'], base_year=base_year)
            if item['date'] and lesson_date is None:
                errors.append('Некоректна дата. Використовуйте РРРР-ММ-ДД або ДД.ММ.РРРР.')
            if lesson_date and not journal.academic_year.start_date <= lesson_date <= journal.academic_year.end_date:
                errors.append('Дата поза навчальним роком цього журналу. Перевірте рік у файлі.')
            item['date'] = lesson_date.isoformat() if lesson_date else ''
            item['date_display'] = lesson_date.strftime('%d.%m.%Y') if lesson_date else 'Без дати (план)'
            item['lesson_type'] = item['lesson_type'].upper() or LessonType.ORDINARY
            if item['lesson_type'] not in LessonType.values:
                errors.append('Невідомий тип уроку. Наприклад: ORDINARY, QUIZ, TEST, THEMATIC.')
            item['record_kind'] = item['record_kind'].upper()
            if item['record_kind'] not in ('', 'LESSON', 'PLAN'):
                errors.append('Вид запису має бути LESSON (урок) або PLAN (план).')
            lesson = lessons_by_id.get(item['lesson_id'])
            plan = plans_by_id.get(item['plan_topic_id'])
            if item['lesson_id'] and lesson is None:
                errors.append('ID уроку відсутній у цьому журналі. Завантажте актуальний файл.')
            if item['plan_topic_id'] and plan is None:
                errors.append('ID теми плану відсутній у цьому журналі. Завантажте актуальний файл.')
            if lesson and plan and lesson.linked_plan_topic_id not in (None, plan.pk):
                errors.append('Урок пов’язаний з іншою темою плану.')
            if not plan and lesson and lesson.linked_plan_topic_id:
                plan = plans_by_id.get(lesson.linked_plan_topic_id)
            if not plan and not item['plan_topic_id']:
                candidates = plans_by_number[item['topic_number']]
                if len(candidates) > 1:
                    errors.append('У плані кілька тем із цим номером. Використайте експорт із ID теми.')
                elif candidates:
                    plan = candidates[0]
            plan_only = target == 'plan_only' or item['record_kind'] == 'PLAN'
            if plan_only:
                lesson = None
            elif not lesson:
                candidates = lessons_by_plan[plan.pk] if plan else []
                if not candidates and lesson_date:
                    candidates = lessons_by_date[lesson_date]
                    if item['day_lesson_number']:
                        candidates = [l for l in candidates if l.lesson_number == item['day_lesson_number']]
                    elif item['lesson_type'] in SUMMARY_TYPES:
                        candidates = [l for l in candidates if l.lesson_type == item['lesson_type']]
                    else:
                        candidates = [l for l in candidates if l.lesson_type not in SUMMARY_TYPES]
                if len(candidates) > 1:
                    errors.append('На цю дату є кілька уроків. Завантажте поточні дані з ID або вкажіть «Номер уроку на дату».')
                elif candidates:
                    lesson = candidates[0]
            if lesson:
                if lesson_date and lesson_date != lesson.date:
                    errors.append('Дата не збігається з наявним уроком. Імпорт тем не переносить уроки; змініть дату окремо.')
                if lesson.nz_schedule_id and ((lesson.lesson_type in SUMMARY_TYPES) != (item['lesson_type'] in SUMMARY_TYPES)
                                              or (lesson.lesson_type in SUMMARY_TYPES and lesson.lesson_type != item['lesson_type'])):
                    errors.append('Тип підсумкової онлайн-колонки змініть у nz.ua та оновіть журнал.')
                if not plan and lesson.linked_plan_topic_id:
                    plan = plans_by_id[lesson.linked_plan_topic_id]
            elif not plan_only and item['record_kind'] == 'LESSON' and not lesson_date:
                errors.append('Новий урок потребує дати. Для майбутньої теми виберіть PLAN.')
            if plan and item['topic_number'] != plan.topic_number:
                errors.append('Номер не збігається з темою плану. Змініть порядок тем у журналі.')
            item.update(matched_lesson_id=lesson.pk if lesson else None, matched_plan_id=plan.pk if plan else None,
                        exists_lesson=bool(lesson), exists_in_plan=bool(plan), errors=errors, warnings=[],
                        create_lesson=bool(not plan_only and not lesson and lesson_date))
            if lesson:
                item['action_label'] = f'Оновити урок {lesson.date:%d.%m.%Y}, №{lesson.lesson_number}'
            elif item['create_lesson']:
                item['action_label'] = 'Створити урок і тему плану'
            else:
                item['action_label'] = 'Оновити тему плану' if plan else 'Додати тему плану'
            preview.append(item)
        numbers = Counter(p['topic_number'] for p in preview)
        lesson_ids = Counter(p['matched_lesson_id'] for p in preview if p['matched_lesson_id'])
        plan_ids = Counter(p['matched_plan_id'] for p in preview if p['matched_plan_id'])
        new_slots = Counter((p['date'], p['day_lesson_number']) for p in preview if p['create_lesson'])
        for item in preview:
            if numbers[item['topic_number']] > 1 or lesson_ids[item['matched_lesson_id']] > 1 or plan_ids[item['matched_plan_id']] > 1:
                item['errors'].append('Кілька рядків змінюють той самий урок або номер теми. Залиште один рядок.')
            if item['create_lesson'] and new_slots[(item['date'], item['day_lesson_number'])] > 1:
                item['errors'].append('Для кількох нових уроків на одну дату вкажіть різні номери уроку на дату.')
            item['is_valid'] = not item['errors']
        return {'preview_data': preview, 'total_rows': len(preview),
                'valid_rows': sum(p['is_valid'] for p in preview), 'error_rows': sum(not p['is_valid'] for p in preview),
                'errors': [f"Рядок {p['row_num']}: {'; '.join(p['errors'])}" for p in preview if p['errors']], 'warnings': []}

    @classmethod
    @transaction.atomic
    def execute_import(cls, preview_data, journal, mode='add_and_update', user=None, target='lessons_and_plan'):
        if mode not in cls.MODES:
            raise ValueError('Невідомий режим імпорту.')
        Journal.objects.select_for_update().get(pk=journal.pk)
        valid_items = [item for item in preview_data if item.get('is_valid', True)]
        current = cls.validate_and_preview(valid_items, journal, target)['preview_data']
        for old, item in zip(valid_items, current):
            if not item['is_valid'] or any(key in old and old[key] != item[key] for key in ('matched_lesson_id', 'matched_plan_id')):
                raise ValueError('Уроки або план змінилися після перегляду. Завантажте файл і перевірте його ще раз.')
        stats = dict.fromkeys(('added_plan', 'updated_plan', 'added_lessons', 'updated_lessons', 'unchanged'), 0)
        stats['skipped'] = len(preview_data) - len(valid_items)
        semesters = list(journal.academic_year.semesters.all())
        for item in current:
            lesson = journal.lessons.filter(pk=item['matched_lesson_id']).first()
            plan = journal.plan_topics.filter(pk=item['matched_plan_id']).first()
            if (mode == 'add_only' and (lesson or plan)) or (mode == 'update_only' and not (lesson or plan)):
                stats['skipped'] += 1
                continue
            lesson_date = datetime.date.fromisoformat(item['date']) if item['date'] else None
            plan_fields = dict(title=item['title'], homework=item['homework'], lesson_type=item['lesson_type'],
                               notes=item['notes'], planned_date=lesson.date if lesson else lesson_date)
            if plan:
                changed = any(getattr(plan, key) != value for key, value in plan_fields.items())
                for key, value in plan_fields.items():
                    setattr(plan, key, value)
                if changed:
                    plan.save()
                    stats['updated_plan'] += 1
            elif mode != 'update_only':
                plan = LessonPlanTopic.objects.create(journal=journal, topic_number=item['topic_number'], **plan_fields)
                stats['added_plan'] += 1
            if lesson:
                fields = {'topic': item['title'], 'homework': item['homework'], 'lesson_type': item['lesson_type'], 'notes': item['notes']}
                changed = any(getattr(lesson, key) != value for key, value in fields.items())
                linked_changed = plan and lesson.linked_plan_topic_id != plan.pk
                for key, value in fields.items():
                    setattr(lesson, key, value)
                if item['homework'] and not (lesson.homework_due_lesson_id or lesson.homework_due_plan_topic_id or lesson.homework_due_date):
                    lesson.homework_for_next_lesson = True
                if plan:
                    lesson.linked_plan_topic = plan
                if changed:
                    lesson.is_synced_to_nz = False
                if changed or linked_changed:
                    lesson.save()
                    stats['updated_lessons'] += 1
                else:
                    stats['unchanged'] += 1
            elif item['create_lesson'] and mode != 'update_only':
                lesson = Lesson.objects.create(journal=journal, date=lesson_date,
                    lesson_number=item['day_lesson_number'] or next_lesson_number(journal, lesson_date),
                    semester=next((s for s in semesters if s.start_date <= lesson_date <= s.end_date), None),
                    topic=item['title'], homework=item['homework'], lesson_type=item['lesson_type'], notes=item['notes'],
                    homework_for_next_lesson=True,
                    linked_plan_topic=plan, created_by=user)
                stats['added_lessons'] += 1
            if plan and lesson and not plan.is_conducted:
                plan.is_conducted = True
                plan.save(update_fields=['is_conducted', 'updated_at'])
        AuditLog.objects.create(school=journal.school, user=user, action=AuditLog.Action.IMPORT,
            target_model='LessonPlanTopic', target_id=str(journal.pk),
            description=f"Імпорт тем: уроків додано {stats['added_lessons']}, оновлено {stats['updated_lessons']}.",
            changes={**stats, 'mode': mode, 'target': target})
        return stats
