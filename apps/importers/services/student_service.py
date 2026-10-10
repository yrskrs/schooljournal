import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.db import transaction
from django.utils import timezone
from apps.curriculum.models import Student, StudentClass, ClassEnrollment
from apps.journals.models import AuditLog


class StudentImportService:
    @staticmethod
    def generate_template(file_format='xlsx'):
        """Генерує файл-шаблон для імпорту учнів із прикладами."""
        headers = [
            "Код/ID учня",
            "Прізвище",
            "Ім'я",
            "По батькові",
            "Клас",
            "Статус (ACTIVE/TRANSFERRED/EXPELLED)"
        ]
        sample_rows = [
            ["STD-001", "Шевченко", "Тарас", "Григорович", "5-А", "ACTIVE"],
            ["STD-002", "Франко", "Іван", "Якович", "5-А", "ACTIVE"],
            ["STD-003", "Косач", "Лариса", "Петрівна", "5-А", "ACTIVE"],
            ["STD-004", "Котляревський", "Іван", "Петрович", "7-Б", "ACTIVE"],
        ]

        if file_format == 'csv':
            output = io.StringIO()
            # UTF-8 with BOM for Ukrainian characters in Excel
            writer = csv.writer(output)
            writer.writerow(headers)
            for row in sample_rows:
                writer.writerow(row)
            return output.getvalue().encode('utf-8-sig'), 'text/csv'
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Учні"

            # Header styling
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

            ws.append(headers)
            for col_idx, cell in enumerate(ws[1], 1):
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for row in sample_rows:
                ws.append(row)

            # Auto fit width
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    @classmethod
    def parse_file(cls, file_obj, filename):
        """Парсить завантажений файл XLSX або CSV у список словників рядків."""
        rows = []
        is_csv = filename.lower().endswith('.csv')

        if is_csv:
            content = file_obj.read()
            # Try utf-8-sig, utf-8, cp1251
            decoded = None
            for enc in ['utf-8-sig', 'utf-8', 'cp1251', 'latin-1']:
                try:
                    decoded = content.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if decoded is None:
                raise ValueError("Не вдалося визначити кодування файлу CSV. Використовуйте UTF-8.")

            reader = csv.reader(io.StringIO(decoded))
            raw_rows = list(reader)
        else:
            wb = openpyxl.load_workbook(file_obj, data_only=True)
            ws = wb.active
            raw_rows = []
            for row in ws.iter_rows(values_only=True):
                # Filter out completely empty rows
                if any(c is not None and str(c).strip() != '' for c in row):
                    raw_rows.append(list(row))

        if not raw_rows:
            return []

        # Визначаємо заголовки
        raw_headers = [str(c or '').strip() for c in raw_rows[0]]
        header_map = cls._map_headers(raw_headers)

        for row_num, raw_row in enumerate(raw_rows[1:], start=2):
            if not any(c is not None and str(c).strip() != '' for c in raw_row):
                continue

            item = {'_row_num': row_num}
            for key, col_idx in header_map.items():
                if col_idx is not None and col_idx < len(raw_row):
                    val = raw_row[col_idx]
                    item[key] = str(val).strip() if val is not None else ""
                else:
                    item[key] = ""
            rows.append(item)

        return rows

    @classmethod
    def _map_headers(cls, headers):
        """Зіставляє заголовки файлу з ключами системи."""
        mapping = {
            'internal_id': None,
            'last_name': None,
            'first_name': None,
            'middle_name': None,
            'class_name': None,
            'status': None,
        }

        for idx, h in enumerate(headers):
            hl = h.lower().strip()
            if any(k in hl for k in ['id', 'код', 'особова', 'номер']):
                mapping['internal_id'] = idx
            elif 'прізвище' in hl:
                mapping['last_name'] = idx
            elif "ім'я" in hl or "імя" in hl:
                mapping['first_name'] = idx
            elif any(k in hl for k in ['батьк', 'по-батькові', 'по батькові']):
                mapping['middle_name'] = idx
            elif 'клас' in hl:
                mapping['class_name'] = idx
            elif 'статус' in hl:
                mapping['status'] = idx

        # Fallback to column position if headers weren't named
        if mapping['last_name'] is None and len(headers) >= 2:
            mapping['last_name'] = 1
        if mapping['first_name'] is None and len(headers) >= 3:
            mapping['first_name'] = 2

        return mapping

    @classmethod
    def validate_and_preview(cls, rows, school, target_class=None):
        """
        Перевіряє обов'язкові поля, виявляє дублікати та формує попередній перегляд.
        target_class: якщо імпорт веде вчитель з конкретного журналу.
        """
        preview_data = []
        errors = []
        warnings = []

        existing_students_by_id = {
            s.internal_id: s for s in Student.objects.filter(school=school) if s.internal_id
        }

        # All students in school for name lookup
        all_school_students = list(Student.objects.filter(school=school))

        for row in rows:
            row_num = row.get('_row_num', '?')
            last_name = row.get('last_name', '').strip()
            first_name = row.get('first_name', '').strip()
            middle_name = row.get('middle_name', '').strip()
            internal_id = row.get('internal_id', '').strip()
            class_name = row.get('class_name', '').strip()
            status_str = row.get('status', '').strip().upper() or 'ACTIVE'

            row_errors = []
            row_warnings = []

            # 1. Обов'язкові поля
            if not last_name:
                row_errors.append("Відсутнє обов'язкове поле: Прізвище")
            if not first_name:
                row_errors.append("Відсутнє обов'язкове поле: Ім'я")

            # 2. Визначення класу
            assigned_class = target_class
            if not assigned_class and class_name:
                matched_classes = list(StudentClass.objects.filter(school=school, name__iexact=class_name))
                if matched_classes:
                    assigned_class = matched_classes[0]
                else:
                    row_warnings.append(f"Клас «{class_name}» не знайдено в системі для даної школи.")
            elif not assigned_class and not class_name:
                row_warnings.append("Клас не вказано. Учень буде створений без прив'язки до класу.")

            # 3. Виявлення дублікатів / збігів
            matched_student = None
            is_duplicate = False

            if internal_id and internal_id in existing_students_by_id:
                matched_student = existing_students_by_id[internal_id]
                is_duplicate = True
                row_warnings.append(f"Знайдено учня з таким ID ({internal_id}): {matched_student.get_full_name_ukr()}")
            else:
                # Перевірка на однакове ПІБ
                same_name = [
                    s for s in all_school_students
                    if s.last_name.lower() == last_name.lower() and s.first_name.lower() == first_name.lower()
                ]
                if same_name:
                    is_duplicate = True
                    matched_student = same_name[0]
                    row_warnings.append(f"Увага: у системі вже є учень зі схожим ПІБ: {matched_student.get_full_name_ukr()} (ID: {matched_student.internal_id or 'немає'})")

            preview_item = {
                'row_num': row_num,
                'internal_id': internal_id,
                'last_name': last_name,
                'first_name': first_name,
                'middle_name': middle_name,
                'class_name': assigned_class.name if assigned_class else class_name,
                'class_id': assigned_class.id if assigned_class else None,
                'status': status_str,
                'is_duplicate': is_duplicate,
                'matched_student_id': matched_student.id if matched_student else None,
                'errors': row_errors,
                'warnings': row_warnings,
                'is_valid': len(row_errors) == 0,
            }
            preview_data.append(preview_item)
            if row_errors:
                errors.append(f"Рядок {row_num}: {'; '.join(row_errors)}")
            if row_warnings:
                warnings.append(f"Рядок {row_num}: {'; '.join(row_warnings)}")

        return {
            'preview_data': preview_data,
            'total_rows': len(preview_data),
            'valid_rows': len([p for p in preview_data if p['is_valid']]),
            'error_rows': len([p for p in preview_data if not p['is_valid']]),
            'errors': errors,
            'warnings': warnings,
        }

    @classmethod
    @transaction.atomic
    def execute_import(cls, preview_data, school, mode='add_and_update', user=None):
        """
        Застосовує імпорт у транзакції Django.
        mode:
          - 'add_and_update': додає нових та оновлює наявні
          - 'add_only': додає лише нових, ігнорує дублікати
          - 'update_only': оновлює лише знайдених
        """
        added_count = 0
        updated_count = 0
        skipped_count = 0
        affected_classes = set()

        for item in preview_data:
            if not item.get('is_valid', True):
                skipped_count += 1
                continue

            internal_id = item.get('internal_id', '').strip()
            last_name = item.get('last_name', '').strip()
            first_name = item.get('first_name', '').strip()
            middle_name = item.get('middle_name', '').strip()
            status_val = item.get('status', 'ACTIVE')
            matched_id = item.get('matched_student_id')
            class_id = item.get('class_id')

            student = None
            if matched_id:
                try:
                    student = Student.objects.get(id=matched_id, school=school)
                except Student.DoesNotExist:
                    student = None

            if student:
                if mode == 'add_only':
                    skipped_count += 1
                    continue
                # Update existing
                old_info = f"{student.get_full_name_ukr()} [{student.status}]"
                student.last_name = last_name
                student.first_name = first_name
                if middle_name:
                    student.middle_name = middle_name
                if internal_id:
                    student.internal_id = internal_id
                if status_val in [s[0] for s in Student.Status.choices]:
                    student.status = status_val
                student.save()
                updated_count += 1
            else:
                if mode == 'update_only':
                    skipped_count += 1
                    continue
                # Create new student
                valid_status = status_val if status_val in [s[0] for s in Student.Status.choices] else Student.Status.ACTIVE
                student = Student.objects.create(
                    school=school,
                    internal_id=internal_id,
                    last_name=last_name,
                    first_name=first_name,
                    middle_name=middle_name,
                    status=valid_status
                )
                added_count += 1

            # Оновлюємо або створюємо зарахування до класу
            if class_id:
                try:
                    student_class = StudentClass.objects.get(id=class_id, school=school)
                    affected_classes.add(student_class.name)
                    # Перевіряємо поточне активне зарахування
                    enrollment, created = ClassEnrollment.objects.get_or_create(
                        student=student,
                        student_class=student_class,
                        defaults={
                            'start_date': timezone.now().date(),
                            'is_active': True,
                            'transfer_note': 'Імпорт зі списку'
                        }
                    )
                    if not created and not enrollment.is_active:
                        enrollment.is_active = True
                        enrollment.end_date = None
                        enrollment.save()
                except StudentClass.DoesNotExist:
                    pass

        # Запис в AuditLog
        classes_str = ", ".join(affected_classes) if affected_classes else "Без класу"
        AuditLog.objects.create(
            school=school,
            user=user,
            action=AuditLog.Action.IMPORT,
            target_model='Student',
            description=f"Імпорт учнів: додано {added_count}, оновлено {updated_count}, пропущено {skipped_count}. Класи: {classes_str}",
            changes={
                'added': added_count,
                'updated': updated_count,
                'skipped': skipped_count,
                'classes': list(affected_classes)
            }
        )

        return {
            'added': added_count,
            'updated': updated_count,
            'skipped': skipped_count,
            'affected_classes': list(affected_classes)
        }
