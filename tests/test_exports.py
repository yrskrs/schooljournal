import io
import datetime
import openpyxl
from django.test import TestCase, Client
from django.urls import reverse
from apps.schools.models import School
from apps.accounts.models import User
from apps.curriculum.models import AcademicYear, StudentClass, Subject, Student, ClassEnrollment
from apps.journals.models import Journal, Lesson, Grade
from apps.importers.services.export_service import ExportService
from apps.importers.services.pdf_service import JournalPDFExportService


class ExportsTestCase(TestCase):
    """
    Тести для сценарію 9:
    - Експорт учнів (XLSX, CSV) з можливістю повторного імпорту.
    - Експорт календарного плану (XLSX, CSV).
    - Експорт журналу в XLSX та читабельний PDF з українським шрифтом.
    """

    def setUp(self):
        self.school = School.objects.create(name="Ліцей «Інтелект»", is_active=True)
        self.year = AcademicYear.objects.create(
            school=self.school, name="2024/2025",
            start_date=datetime.date(2024, 9, 1), end_date=datetime.date(2025, 6, 30),
            is_current=True
        )
        self.student_class = StudentClass.objects.create(
            school=self.school, academic_year=self.year, name="10-А", grade_level=10
        )
        self.subj = Subject.objects.create(school=self.school, name="Українська мова")
        self.teacher = User.objects.create_user(
            username="teacher_export", password="password123", role=User.Role.TEACHER, school=self.school,
            last_name="Тищенко", first_name="Світлана"
        )
        self.journal = Journal.objects.create(
            school=self.school, academic_year=self.year, student_class=self.student_class,
            subject=self.subj, teacher=self.teacher
        )
        self.student = Student.objects.create(
            school=self.school, internal_id="STD-101",
            last_name="Гончар", first_name="Олесь", middle_name="Терентійович"
        )
        ClassEnrollment.objects.create(
            student=self.student, student_class=self.student_class, start_date=datetime.date(2024, 9, 1)
        )
        self.lesson = Lesson.objects.create(
            journal=self.journal, date=datetime.date(2024, 9, 4), lesson_number=1, topic="Стилі мовлення"
        )
        Grade.objects.create(lesson=self.lesson, student=self.student, value="12")

    def test_export_students_xlsx_and_csv(self):
        """Експорт учнів повертає коректний файл, придатний для повторного імпорту."""
        students = [self.student]

        # 1. XLSX
        xlsx_bytes, mime = ExportService.export_students(students, student_class=self.student_class, file_format='xlsx')
        self.assertIn('spreadsheetml', mime)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        self.assertEqual(ws.cell(row=2, column=2).value, "Гончар")
        self.assertEqual(ws.cell(row=2, column=3).value, "Олесь")
        self.assertEqual(ws.cell(row=2, column=5).value, "10-А")

        # 2. CSV
        csv_bytes, mime_csv = ExportService.export_students(students, student_class=self.student_class, file_format='csv')
        self.assertEqual(mime_csv, 'text/csv')
        csv_text = csv_bytes.decode('utf-8-sig')
        self.assertIn("Гончар", csv_text)
        self.assertIn("STD-101", csv_text)

    def test_export_topics_xlsx_and_csv(self):
        """Експорт тем уроків повертає коректний файл."""
        xlsx_bytes, _ = ExportService.export_topics(journal=self.journal, file_format='xlsx')
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        self.assertEqual(ws.cell(row=2, column=2).value, "Стилі мовлення")

    def test_export_journal_xlsx(self):
        """Експорт журналу в XLSX містить метадані школи, учнів, уроків та оцінок."""
        xlsx_bytes, _ = ExportService.export_journal_xlsx(self.journal)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        # Метадані
        self.assertIn("Ліцей «Інтелект»", ws['A1'].value)
        self.assertIn("10-А", ws['A2'].value)
        # Учень та оцінка
        cell_student = ws.cell(row=7, column=2).value
        self.assertIn("Гончар Олесь", cell_student)
        grade_val = ws.cell(row=7, column=3).value
        self.assertEqual(grade_val, "12")

    def test_export_journal_pdf(self):
        """Експорт журналу в PDF повертає валідний PDF документ з українським текстом."""
        pdf_bytes = JournalPDFExportService.generate_journal_pdf(self.journal)
        self.assertTrue(pdf_bytes.startswith(b'%PDF-'))
        self.assertGreater(len(pdf_bytes), 1000)

    def test_export_endpoints_response(self):
        """Ендпоінти завантаження повертають HTTP 200 та файлові заголовки Content-Disposition."""
        client = Client()
        client.login(username="teacher_export", password="password123")

        # XLSX
        url_xlsx = reverse('portal:export_journal_xlsx', kwargs={'journal_id': self.journal.id})
        res_xlsx = client.get(url_xlsx)
        self.assertEqual(res_xlsx.status_code, 200)
        self.assertIn('attachment', res_xlsx['Content-Disposition'])

        # PDF
        url_pdf = reverse('portal:export_journal_pdf', kwargs={'journal_id': self.journal.id})
        res_pdf = client.get(url_pdf)
        self.assertEqual(res_pdf.status_code, 200)
        self.assertEqual(res_pdf['Content-Type'], 'application/pdf')
        self.assertIn('attachment', res_pdf['Content-Disposition'])
